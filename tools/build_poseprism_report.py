from pathlib import Path
import math
import random

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(r"D:\02-学习\项目实训\4")
OUT = ROOT / "deliverables"
ASSETS = ROOT / "frontend" / "src" / "assets"
PUBLIC = ROOT / "frontend" / "public"
LOGO = Path(r"D:\02-学习\项目实训\logo\938b2dd0fea680c4f870aa5bca309594.png")

OUT.mkdir(exist_ok=True)

W, H = 13.333, 7.5

COLORS = {
    "navy": "07111F",
    "ink": "111827",
    "muted": "667085",
    "soft": "F6F8FF",
    "card": "FFFFFF",
    "line": "E6ECF8",
    "violet": "6C3BFF",
    "cyan": "43B7FF",
    "mint": "44D989",
    "coral": "FF5B57",
    "amber": "FFB84D",
}


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def make_shadow():
    return {"type": "outer", "color": "000000", "opacity": 0.12, "blur": 2, "offset": 1, "angle": 45}


def generate_visual_assets():
    bg = Image.new("RGB", (1600, 900), "#07111f")
    px = bg.load()
    for y in range(bg.height):
        for x in range(bg.width):
            t = (x / bg.width) * 0.55 + (y / bg.height) * 0.45
            r = int(7 + 22 * t)
            g = int(17 + 28 * t)
            b = int(31 + 88 * t)
            px[x, y] = (r, g, b)
    overlay = Image.new("RGBA", bg.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(overlay)
    for cx, cy, color in [(1180, 160, (67, 183, 255, 80)), (1350, 680, (68, 217, 137, 60)), (260, 710, (108, 59, 255, 80))]:
        for radius in range(360, 0, -8):
            alpha = int(color[3] * (radius / 360) ** 2)
            d.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=(*color[:3], alpha))
    bg = Image.alpha_composite(bg.convert("RGBA"), overlay).filter(ImageFilter.GaussianBlur(1))
    bg.save(OUT / "deck_bg.png")

    mock = Image.new("RGB", (1300, 720), "#050914")
    d = ImageDraw.Draw(mock, "RGBA")
    d.rounded_rectangle((40, 40, 1260, 680), radius=42, fill=(8, 12, 28, 255), outline=(55, 82, 130, 150), width=2)
    d.rounded_rectangle((80, 70, 440, 134), radius=28, fill=(20, 26, 52, 240))
    d.ellipse((120, 94, 142, 116), fill=(67, 183, 255, 255))
    d.text((155, 88), "User", fill=(245, 248, 255, 255))
    d.ellipse((295, 94, 317, 116), fill=(68, 217, 137, 255))
    d.text((330, 88), "Std", fill=(245, 248, 255, 255))
    d.line((650, 64, 650, 670), fill=(95, 119, 170, 110), width=2)
    random.seed(7)
    for side, tint, cx in [("left", (67, 183, 255), 350), ("right", (68, 217, 137), 950)]:
        joints = [(cx, 185), (cx-55, 295), (cx+55, 295), (cx-45, 430), (cx+45, 430), (cx-78, 565), (cx+78, 565),
                  (cx-120, 365), (cx+120, 365), (cx-145, 510), (cx+145, 510)]
        bones = [(0,1),(0,2),(1,2),(1,3),(2,4),(3,4),(3,5),(4,6),(1,7),(2,8),(7,9),(8,10)]
        for a, b in bones:
            d.line((joints[a], joints[b]), fill=(*tint, 95), width=5)
        for x, y in joints:
            d.ellipse((x-10, y-10, x+10, y+10), fill=(*tint, 220))
        for _ in range(1100):
            j = random.choice(joints)
            x = int(random.gauss(j[0], 42))
            y = int(random.gauss(j[1], 42))
            if 70 < x < 1230 and 90 < y < 650:
                d.ellipse((x, y, x+3, y+3), fill=(*tint, random.randint(80, 210)))
    mock.save(OUT / "mock_3d_compare.png")

    flow = Image.new("RGBA", (1200, 520), (255, 255, 255, 0))
    d = ImageDraw.Draw(flow)
    boxes = [
        ("小程序/摄像头", "采集训练画面"),
        ("姿态识别", "提取人体关键点"),
        ("动作分析", "计数、评分、纠错"),
        ("Web 可视化", "3D 回放与报告"),
    ]
    colors = ["#43B7FF", "#6C3BFF", "#44D989", "#FFB84D"]
    for i, (title, sub) in enumerate(boxes):
        x = 25 + i * 290
        d.rounded_rectangle((x, 140, x + 245, 330), radius=24, fill=(255, 255, 255, 255), outline=(216, 226, 246, 255), width=3)
        d.ellipse((x + 24, 176, x + 74, 226), fill=colors[i])
        d.text((x + 92, 174), title, fill="#111827")
        d.text((x + 92, 218), sub, fill="#667085")
        if i < 3:
            d.line((x + 250, 235, x + 285, 235), fill=(108, 59, 255, 200), width=5)
            d.polygon([(x + 285, 235), (x + 270, 225), (x + 270, 245)], fill=(108, 59, 255, 220))
    flow.save(OUT / "flow_arch.png")


generate_visual_assets()


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)


def blank():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return slide


def set_bg(slide, color="F6F8FF"):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_logo(slide, x=11.75, y=0.28, w=0.62):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(x), Inches(y), width=Inches(w), height=Inches(w))


def add_text(slide, text, x, y, w, h, size=18, color="111827", bold=False, align="left", font="Microsoft YaHei", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def title(slide, text, sub=None, dark=False):
    add_text(slide, text, 0.7, 0.38, 8.8, 0.55, 26 if len(text) > 15 else 30, "FFFFFF" if dark else "111827", True)
    if sub:
        add_text(slide, sub, 0.72, 0.96, 8.3, 0.32, 10.5, "B9C7E6" if dark else "667085")
    add_logo(slide)


def footer(slide, page):
    add_text(slide, f"{page:02d} / 姿态棱镜课程汇报", 10.7, 7.12, 2.0, 0.2, 8.5, "98A2B3", align="right")


def rect(slide, x, y, w, h, fill="FFFFFF", line="E6ECF8", radius=True, shadow=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    if shadow:
        try:
            s.shadow.inherit = False
            s.shadow.blur_radius = Pt(6)
            s.shadow.distance = Pt(2)
            s.shadow.angle = 45
            s.shadow.fore_color.rgb = rgb("000000")
            s.shadow.transparency = 85
        except Exception:
            pass
    return s


def pill(slide, text, x, y, fill, color="FFFFFF", w=None):
    width = w or max(1.0, 0.18 * len(text) + 0.45)
    rect(slide, x, y, width, 0.38, fill, None, True)
    add_text(slide, text, x, y + 0.09, width, 0.2, 10, color, True, align="center")


def bullet_lines(slide, items, x, y, w, h, size=14, color="344054"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(8)
    return box


def image(slide, path, x, y, w, h):
    if Path(path).exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def full_image_bg(slide, path):
    if Path(path).exists():
        slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=Inches(W), height=Inches(H))


def add_card(slide, heading, body, x, y, w, h, accent="6C3BFF", icon=None):
    rect(slide, x, y, w, h, "FFFFFF", "E7ECFA", True, True)
    rect(slide, x + 0.18, y + 0.22, 0.12, h - 0.44, accent, None, False)
    add_text(slide, heading, x + 0.45, y + 0.25, w - 0.65, 0.3, 14.5, "111827", True)
    add_text(slide, body, x + 0.45, y + 0.68, w - 0.65, h - 0.85, 10.8, "667085")


def cover():
    s = blank()
    full_image_bg(s, OUT / "deck_bg.png")
    image(s, LOGO, 0.72, 0.52, 0.78, 0.78)
    add_text(s, "姿态棱镜", 0.78, 1.72, 5.5, 0.75, 42, "FFFFFF", True)
    add_text(s, "AI 运动姿态评估与纠错系统", 0.82, 2.52, 6.8, 0.45, 22, "D7E8FF", True)
    add_text(s, "让每一次训练都能被看见、被分析、被改进", 0.84, 3.1, 5.8, 0.38, 15, "AFC5E9")
    pill(s, "课程项目汇报", 0.86, 4.0, "6C3BFF")
    pill(s, "3D 回放  ·  标准动作对比  ·  真实纠错", 2.15, 4.0, "0F2748", "D7E8FF", 3.4)
    image(s, OUT / "mock_3d_compare.png", 7.0, 1.0, 5.6, 3.1)
    add_text(s, "团队名称 / 汇报人 / 日期", 0.86, 6.82, 5.0, 0.28, 11, "B9C7E6")


def toc():
    s = blank()
    set_bg(s)
    title(s, "目录", "从团队、背景、架构到功能创新与未来规划")
    items = [("01", "团队成员介绍"), ("02", "项目背景"), ("03", "技术架构解析"), ("04", "项目功能与创新点"), ("05", "未来发展展望")]
    for i, (num, text) in enumerate(items):
        x = 0.9 + (i % 3) * 4.05
        y = 1.55 + (i // 3) * 2.2
        rect(s, x, y, 3.4, 1.55, "FFFFFF", "E2E9FA", True, True)
        add_text(s, num, x + 0.28, y + 0.24, 0.68, 0.38, 20, "6C3BFF", True)
        add_text(s, text, x + 0.28, y + 0.82, 2.7, 0.34, 17, "111827", True)
    image(s, ASSETS / "start-tablet-replay.png", 9.1, 4.35, 2.1, 2.1)
    footer(s, 2)


def team():
    s = blank()
    set_bg(s)
    title(s, "团队成员介绍", "成员姓名可根据实际团队信息替换")
    roles = [
        ("成员 1", "项目统筹 / 汇报", "需求梳理、演示流程、汇报材料整合"),
        ("成员 2", "前端与交互", "Web 端页面、3D 回放、可视化体验"),
        ("成员 3", "后端与接口", "FastAPI、训练记录、回放数据管理"),
        ("成员 4", "姿态算法", "关键点识别、计数、评分与纠错"),
        ("成员 5", "小程序端", "训练采集、骨架数据上传、移动端体验"),
    ]
    for i, (name, role, work) in enumerate(roles):
        x = 0.68 + (i % 3) * 4.15
        y = 1.45 + (i // 3) * 2.25
        add_card(s, role, work, x, y, 3.55, 1.65, ["43B7FF", "6C3BFF", "44D989", "FFB84D", "FF5B57"][i])
        add_text(s, name, x + 2.55, y + 0.22, 0.72, 0.26, 10.5, "667085", True, align="right")
    footer(s, 3)


def background_data():
    s = blank()
    set_bg(s)
    title(s, "项目背景：运动热潮下的新问题", "大家更愿意运动，但更难知道自己练得对不对")
    stats = [
        ("18亿", "全球约 18 亿成年人身体活动不足", "WHO 2024"),
        ("80%+", "全球超过 80% 青少年运动不足", "WHO 2022"),
        ("5.5亿", "中国经常参加体育锻炼人数约 5.5 亿", "China Briefing 2023"),
    ]
    for i, (num, label, src) in enumerate(stats):
        x = 0.75 + i * 4.15
        rect(s, x, 1.55, 3.55, 2.25, "FFFFFF", "E6ECF8", True, True)
        add_text(s, num, x + 0.3, 1.9, 2.9, 0.72, 42, ["6C3BFF", "43B7FF", "44D989"][i], True)
        add_text(s, label, x + 0.34, 2.72, 2.85, 0.5, 13, "111827", True)
        add_text(s, src, x + 0.34, 3.35, 2.0, 0.22, 9.5, "98A2B3")
    rect(s, 0.82, 4.35, 11.6, 1.35, "07111F", None, True)
    add_text(s, "痛点不是“不想动”，而是“不知道动作是否标准”", 1.2, 4.75, 9.8, 0.38, 22, "FFFFFF", True)
    add_text(s, "普通用户缺少低成本、可复盘、能看懂的训练指导工具。", 1.22, 5.24, 9.6, 0.28, 12.5, "B9C7E6")
    footer(s, 4)


def pain_video():
    s = blank()
    set_bg(s, "FFFFFF")
    title(s, "真实场景视频：不请教练，也能看懂动作", "建议把你们录制的男生使用姿态棱镜视频放在这一页")
    rect(s, 0.78, 1.35, 7.25, 4.75, "050914", "15233B", True, True)
    add_text(s, "放置演示视频", 2.82, 3.1, 2.95, 0.5, 28, "FFFFFF", True, align="center")
    add_text(s, "男生不想花钱请教练，使用 Web 端 3D 回放复盘训练", 1.52, 3.72, 5.8, 0.32, 13, "B9C7E6", align="center")
    tri = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, Inches(4.0), Inches(2.35), Inches(0.65), Inches(0.65))
    tri.rotation = 90
    tri.fill.solid(); tri.fill.fore_color.rgb = rgb("43B7FF"); tri.line.fill.background()
    add_card(s, "为什么放在背景章节？", "视频先讲清楚真实需求，再进入系统方案：普通人想练、又担心练错，需要一个低成本的 AI 训练助手。", 8.55, 1.48, 3.9, 1.7, "6C3BFF")
    add_card(s, "讲解重点", "强调“3D 回放做得好”：用户能回看完整训练、查看单次动作、和标准动作同屏对比。", 8.55, 3.55, 3.9, 1.7, "44D989")
    footer(s, 5)


def positioning():
    s = blank()
    set_bg(s)
    title(s, "项目定位：每个人都能使用的 AI 运动教练", "姿态棱镜 = 姿态识别 + 3D 回放 + 真实纠错 + 训练记录")
    image(s, ASSETS / "home-hero-posture.png", 8.15, 1.1, 4.0, 2.25)
    items = [
        ("看得见动作", "摄像头捕捉人体关键点，生成骨架数据"),
        ("看得懂问题", "把专业动作问题转成易理解的建议"),
        ("看得到进步", "保存训练记录、错误次数和历史表现"),
        ("能复盘动作", "完整训练与单次动作都可 3D 回放"),
    ]
    for i, (h, b) in enumerate(items):
        add_card(s, h, b, 0.78 + (i % 2) * 3.75, 1.35 + (i // 2) * 1.85, 3.35, 1.38, ["43B7FF", "6C3BFF", "44D989", "FFB84D"][i])
    rect(s, 8.15, 3.8, 4.0, 1.35, "07111F", None, True)
    add_text(s, "适用场景", 8.5, 4.08, 3.0, 0.28, 15, "FFFFFF", True)
    add_text(s, "居家健身 / 体育教学 / 康复辅助 / 健身房自助训练", 8.5, 4.52, 3.1, 0.36, 11.5, "B9C7E6")
    footer(s, 6)


def architecture():
    s = blank()
    set_bg(s, "FFFFFF")
    title(s, "技术架构解析：从画面到建议", "一条完整链路，把训练画面变成可理解的反馈")
    steps = [
        ("小程序 / 摄像头", "采集训练画面", "43B7FF"),
        ("姿态关键点识别", "提取人体骨架", "6C3BFF"),
        ("后端动作分析", "计数、评分、纠错", "44D989"),
        ("Web 可视化", "3D 回放与报告", "FFB84D"),
    ]
    for i, (h, b, c) in enumerate(steps):
        x = 0.78 + i * 3.15
        rect(s, x, 2.08, 2.55, 1.65, "FFFFFF", "DCE7FA", True, True)
        circle = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.26), Inches(2.44), Inches(0.56), Inches(0.56))
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(c)
        circle.line.fill.background()
        add_text(s, h, x + 0.94, 2.34, 1.35, 0.28, 13, "111827", True)
        add_text(s, b, x + 0.94, 2.82, 1.3, 0.28, 10.5, "667085")
        if i < 3:
            add_text(s, "→", x + 2.68, 2.56, 0.45, 0.35, 26, "6C3BFF", True, align="center")
    rect(s, 1.0, 5.08, 11.25, 0.92, "F6F8FF", "E6ECF8", True)
    add_text(s, "用户端负责采集，后端负责分析，Web 端负责把结果变成直观的 3D 回放与报告。", 1.45, 5.42, 10.4, 0.25, 13, "344054", align="center")
    footer(s, 7)


def core_tech():
    s = blank()
    set_bg(s)
    title(s, "核心技术模块", "少讲术语，多讲每一层解决什么问题")
    modules = [
        ("Vue 3 + Vite", "Web 工作台与训练记录展示", "43B7FF"),
        ("FastAPI", "接口服务、训练数据与报告管理", "6C3BFF"),
        ("MediaPipe", "人体关键点识别与姿态提取", "44D989"),
        ("Three.js", "3D 骨架回放和标准动作对比", "FFB84D"),
        ("小程序端", "移动端采集、上传骨架数据", "FF5B57"),
        ("SQLite / Storage", "训练记录、模板与回放帧保存", "43B7FF"),
    ]
    for i, (h, b, c) in enumerate(modules):
        x = 0.72 + (i % 3) * 4.12
        y = 1.35 + (i // 3) * 2.0
        add_card(s, h, b, x, y, 3.58, 1.42, c)
    footer(s, 8)


def function_capture():
    s = blank()
    set_bg(s, "FFFFFF")
    title(s, "功能一：训练采集与动作识别", "不用穿戴设备，只需要摄像头或小程序采集")
    image(s, ASSETS / "start-phone-camera.png", 0.9, 1.4, 2.0, 2.0)
    image(s, ASSETS / "start-step-phone.png", 3.2, 1.55, 1.45, 1.45)
    image(s, ASSETS / "start-step-cloud.png", 5.2, 1.55, 1.45, 1.45)
    image(s, ASSETS / "start-step-laptop.png", 7.2, 1.55, 1.45, 1.45)
    add_text(s, "采集", 3.1, 3.15, 1.6, 0.28, 15, "111827", True, align="center")
    add_text(s, "分析", 5.1, 3.15, 1.6, 0.28, 15, "111827", True, align="center")
    add_text(s, "展示", 7.1, 3.15, 1.6, 0.28, 15, "111827", True, align="center")
    bullet_lines(s, ["自动识别人体关键点", "判断动作阶段并自动计数", "保存训练过程中的骨架数据", "为后续 3D 回放与纠错提供依据"], 9.0, 1.55, 3.2, 2.8)
    rect(s, 0.95, 4.72, 11.35, 1.08, "F6F8FF", "E6ECF8", True)
    add_text(s, "一句话：把普通训练画面转化为可以计算、可以回看的运动数据。", 1.3, 5.08, 10.5, 0.3, 16, "6C3BFF", True, align="center")
    footer(s, 9)


def function_replay():
    s = blank()
    set_bg(s)
    title(s, "功能二：Web 端 3D 动作回放", "训练结束后，不只留下分数，还能回看动作过程")
    image(s, OUT / "mock_3d_compare.png", 0.72, 1.25, 7.15, 3.95)
    add_card(s, "完整训练", "查看小程序传回来的整段骨架数据，适合观察整体节奏。", 8.3, 1.32, 3.8, 1.18, "43B7FF")
    add_card(s, "单次动作", "在“第 1 次、第 2 次……”之间切换，复盘每一次动作。", 8.3, 2.82, 3.8, 1.18, "6C3BFF")
    add_card(s, "播放控制", "支持播放、暂停、调速、进度拖动和视角切换。", 8.3, 4.32, 3.8, 1.18, "44D989")
    footer(s, 10)


def function_compare():
    s = blank()
    set_bg(s, "FFFFFF")
    title(s, "功能三：标准动作 3D 对比", "左侧看自己，右侧看标准，差异更直观")
    image(s, OUT / "mock_3d_compare.png", 0.7, 1.22, 7.1, 3.92)
    rect(s, 8.25, 1.28, 3.9, 1.1, "07111F", None, True)
    add_text(s, "对比价值", 8.58, 1.55, 3.0, 0.25, 15, "FFFFFF", True)
    bullet_lines(s, ["帮助用户理解动作幅度差异", "减少“文字建议看不懂”的问题", "支持单次动作切换对比", "适合课堂演示和训练复盘"], 8.35, 2.7, 3.8, 2.1, 12.5)
    footer(s, 11)


def correction_records():
    s = blank()
    set_bg(s)
    title(s, "功能四：真实纠错与训练记录", "只展示后端真实返回的问题，不使用静态假提示")
    add_card(s, "真实问题接入", "右侧纠错只读取 rep.issues：没有真实 issue 就不显示该次数。", 0.8, 1.35, 3.7, 1.35, "6C3BFF")
    add_card(s, "单次状态动态化", "下面卡片的“正常/需调整”也由真实问题决定，不再只看分数。", 4.8, 1.35, 3.7, 1.35, "43B7FF")
    add_card(s, "历史复盘", "训练记录保存完整骨架、次数、评分与问题，方便后续查看。", 8.8, 1.35, 3.7, 1.35, "44D989")
    rect(s, 1.0, 3.35, 11.2, 2.0, "FFFFFF", "E6ECF8", True, True)
    add_text(s, "页面效果", 1.38, 3.72, 1.8, 0.3, 16, "111827", True)
    add_text(s, "有真实 issue：显示“需调整”和对应建议；无真实 issue：显示“正常”，右侧不造假纠错卡片。", 1.38, 4.24, 9.8, 0.45, 18, "344054", True)
    footer(s, 12)


def innovation():
    s = blank()
    set_bg(s, "FFFFFF")
    title(s, "项目功能与创新点", "低成本、可复盘、可理解，是姿态棱镜的核心体验")
    points = [
        ("无需穿戴设备", "普通摄像头即可完成动作捕捉"),
        ("3D 回放复盘", "完整训练和单次动作都能回看"),
        ("标准动作对比", "让动作差异直接可见"),
        ("多端联动", "小程序采集，Web 端展示分析"),
        ("真实数据驱动", "纠错提示来自后端分析结果"),
        ("大众友好", "尽量把专业动作问题翻译成易懂建议"),
    ]
    for i, (h, b) in enumerate(points):
        x = 0.8 + (i % 3) * 4.05
        y = 1.35 + (i // 3) * 1.85
        add_card(s, h, b, x, y, 3.55, 1.28, ["43B7FF", "6C3BFF", "44D989", "FFB84D", "FF5B57", "43B7FF"][i])
    footer(s, 13)


def value():
    s = blank()
    set_bg(s)
    title(s, "项目价值：让专业指导更普惠", "不同群体都能从系统中获得清晰反馈")
    groups = [
        ("普通用户", "降低居家运动风险，知道自己练得对不对"),
        ("体育老师", "快速查看学生动作问题，提升教学效率"),
        ("健身教练", "辅助复盘训练表现，减少重复讲解"),
        ("平台方", "沉淀运动数据，扩展个性化训练服务"),
    ]
    for i, (h, b) in enumerate(groups):
        add_card(s, h, b, 0.8 + (i % 2) * 5.95, 1.4 + (i // 2) * 2.0, 5.25, 1.35, ["6C3BFF", "43B7FF", "44D989", "FFB84D"][i])
    image(s, ASSETS / "dashboard-select-coach.png", 8.6, 5.1, 2.9, 1.35)
    footer(s, 14)


def future():
    s = blank()
    set_bg(s, "FFFFFF")
    title(s, "未来发展展望", "从动作识别工具，走向个人 AI 运动教练")
    roadmap = [
        ("近期", "支持更多动作\n完善真实纠错文案\n优化 3D 对比体验"),
        ("中期", "训练报告自动生成\n个性化训练计划\n接入更多健康数据"),
        ("长期", "学校体育与健身房场景\nAI 私教式陪伴\n多人训练数据管理"),
    ]
    for i, (h, b) in enumerate(roadmap):
        x = 0.95 + i * 4.15
        rect(s, x, 1.65, 3.45, 3.3, "FFFFFF", "E6ECF8", True, True)
        pill(s, h, x + 0.35, 2.0, ["43B7FF", "6C3BFF", "44D989"][i], w=1.2)
        add_text(s, b, x + 0.35, 2.75, 2.75, 1.4, 15, "111827", True)
    add_text(s, "目标：让每个人都能用更低成本获得更科学、更安全的运动指导。", 1.35, 5.72, 10.6, 0.4, 18, "6C3BFF", True, align="center")
    footer(s, 15)


def closing_sources():
    s = blank()
    full_image_bg(s, OUT / "deck_bg.png")
    image(s, LOGO, 0.78, 0.62, 0.72, 0.72)
    add_text(s, "谢谢观看", 0.85, 1.78, 4.6, 0.7, 40, "FFFFFF", True)
    add_text(s, "姿态棱镜：让运动更科学，让每个人都能安全地练起来", 0.88, 2.55, 7.0, 0.35, 16, "D7E8FF")
    rect(s, 0.9, 4.25, 11.6, 1.55, "0B1B33", "183355", True)
    add_text(s, "数据来源", 1.2, 4.55, 1.5, 0.25, 13, "FFFFFF", True)
    add_text(s, "WHO 2024 身体活动不足报告；WHO 2022 全球身体活动状况报告；China Briefing 体育健身行业数据；ACSM 健身趋势报告。", 2.55, 4.55, 9.25, 0.48, 11, "B9C7E6")
    add_text(s, "演示建议：最后切到系统实机页面，展示 3D 回放与标准动作对比。", 1.2, 5.23, 9.6, 0.25, 11, "43B7FF", True)


slides = [
    cover,
    toc,
    team,
    background_data,
    pain_video,
    positioning,
    architecture,
    core_tech,
    function_capture,
    function_replay,
    function_compare,
    correction_records,
    innovation,
    value,
    future,
    closing_sources,
]

for fn in slides:
    fn()

prs.core_properties.title = "姿态棱镜课程项目汇报"
prs.core_properties.subject = "AI 运动姿态评估与纠错系统"
prs.core_properties.author = "姿态棱镜项目组"
prs.save(OUT / "姿态棱镜_课程项目汇报.pptx")
print(OUT / "姿态棱镜_课程项目汇报.pptx")
