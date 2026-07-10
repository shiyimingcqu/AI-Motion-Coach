from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


ROOT = Path(r"D:\02-学习\项目实训\4")
OUT = ROOT / "deliverables"
PPT_IN = OUT / "姿态棱镜_课程项目汇报.pptx"
PPT_OUT = OUT / "姿态棱镜_课程项目汇报_新增三页.pptx"
ASSET_DIR = OUT / "added_feature_slide_assets"
SRC_ASSETS = ROOT / "frontend" / "src" / "assets"

W, H = 1920, 1080
FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")

INK = (13, 28, 61)
MUTED = (82, 101, 137)
BLUE = (67, 183, 255)
VIOLET = (108, 59, 255)
GREEN = (68, 217, 137)
AMBER = (255, 184, 77)
CORAL = (255, 91, 87)


def f(size, bold=False):
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT), size)


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def bg():
    img = Image.new("RGB", (W, H), "#F6F8FF").convert("RGBA")
    d = ImageDraw.Draw(img, "RGBA")
    for cx, cy, r, color in [
        (1680, -80, 360, (215, 229, 255, 180)),
        (-120, 880, 360, (223, 245, 255, 160)),
        (940, 1130, 520, (226, 232, 255, 150)),
    ]:
        d.ellipse((cx - r, cy - r, cx + r, cy + r), fill=color)
    for y in [910, 940, 975]:
        d.arc((-260, y - 300, 2200, y + 300), 190, 350, fill=(204, 219, 255, 125), width=3)
    return img


def shadow_card(img, box, radius=36, fill=(255, 255, 255, 235)):
    x1, y1, x2, y2 = box
    shadow = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    sd = ImageDraw.Draw(shadow)
    sd.rounded_rectangle((x1 + 10, y1 + 14, x2 + 10, y2 + 14), radius=radius, fill=(70, 95, 150, 30))
    shadow = shadow.filter(ImageFilter.GaussianBlur(16))
    img.alpha_composite(shadow)
    d = ImageDraw.Draw(img, "RGBA")
    d.rounded_rectangle(box, radius=radius, fill=fill, outline=(222, 231, 248, 255), width=2)
    return d


def text(d, xy, s, size=30, color=INK, bold=False, anchor=None):
    d.text(xy, s, font=f(size, bold), fill=color, anchor=anchor)


def wrap(d, s, font, width):
    lines, cur = [], ""
    for ch in s:
        trial = cur + ch
        if d.textlength(trial, font=font) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = ch
    if cur:
        lines.append(cur)
    return lines


def multiline(d, xy, s, size=28, color=MUTED, bold=False, width=560, gap=10):
    x, y = xy
    font = f(size, bold)
    for line in wrap(d, s, font, width):
        d.text((x, y), line, font=font, fill=color)
        y += size + gap
    return y


def title(img, main, sub, page):
    d = ImageDraw.Draw(img, "RGBA")
    text(d, (86, 58), main, 58, INK, True)
    text(d, (90, 132), sub, 26, MUTED)
    text(d, (1780, 1025), f"{page:02d} / 姿态棱镜课程汇报", 20, (150, 162, 185), False, "ra")


def pill(d, x, y, label, color=VIOLET, bg_fill=(243, 238, 255, 240)):
    font = f(26, True)
    tw = int(d.textlength(label, font=font))
    d.rounded_rectangle((x, y, x + tw + 54, y + 48), radius=24, fill=bg_fill)
    d.text((x + 27, y + 9), label, font=font, fill=color)
    return x + tw + 54


def load_asset(name, size):
    path = SRC_ASSETS / name
    if not path.exists():
        return None
    im = Image.open(path).convert("RGBA")
    im.thumbnail(size, Image.Resampling.LANCZOS)
    return im


def paste_center(img, asset, box):
    if asset is None:
        return
    x1, y1, x2, y2 = box
    x = x1 + (x2 - x1 - asset.width) // 2
    y = y1 + (y2 - y1 - asset.height) // 2
    img.alpha_composite(asset, (x, y))


def slide_action_library():
    img = bg()
    title(img, "新增功能：动作库与动作详情跟练", "用户先选动作，再看详情，最后进入跟练或训练采集", 13)
    d = shadow_card(img, (110, 190, 950, 835))
    text(d, (170, 245), "动作库", 38, INK, True)
    d.rounded_rectangle((170, 310, 880, 370), radius=30, fill=(244, 247, 255, 255), outline=(228, 236, 249, 255))
    text(d, (210, 326), "搜索动作 / 按部位筛选", 24, MUTED)
    for i, name in enumerate(["深蹲", "开合跳", "平板支撑", "弓步蹲"]):
        x = 170 + (i % 2) * 360
        y = 425 + (i // 2) * 150
        d.rounded_rectangle((x, y, x + 310, y + 110), radius=28, fill=(248, 251, 255, 255), outline=(226, 236, 251, 255))
        color = [BLUE, VIOLET, GREEN, AMBER][i]
        d.ellipse((x + 26, y + 28, x + 82, y + 84), fill=color)
        text(d, (x + 110, y + 26), name, 30, INK, True)
        text(d, (x + 110, y + 68), "适合日常训练", 21, MUTED)
    x = pill(d, 170, 710, "按动作分类", BLUE, (232, 244, 255, 240))
    pill(d, x + 18, 710, "快速选择", VIOLET)

    d = shadow_card(img, (1010, 190, 1810, 835))
    text(d, (1070, 245), "动作详情 + 跟练", 38, INK, True)
    d.rounded_rectangle((1070, 315, 1375, 520), radius=30, fill=(9, 17, 36, 245))
    paste_center(img, load_asset("home-exercise-bird-dog.png", (270, 210)), (1395, 286, 1710, 536))
    x = pill(d, 1070, 565, "难度", BLUE, (232, 244, 255, 240))
    x = pill(d, x + 16, 565, "目标部位", VIOLET)
    pill(d, x + 16, 565, "动作要点", GREEN, (230, 250, 241, 240))
    d.rounded_rectangle((1070, 650, 1305, 718), radius=34, fill=VIOLET)
    text(d, (1188, 670), "开始跟练", 28, (255, 255, 255), True, "ma")
    multiline(d, (1070, 748), "动作详情页讲清楚训练目标、动作说明和注意事项，跟练入口直接承接摄像头采集。", 25, MUTED, width=600)
    d = shadow_card(img, (168, 900, 1752, 985), 34, (255, 255, 255, 230))
    text(d, (220, 925), "页面价值：把“找动作、学动作、开始练”连在一起，降低普通用户第一次使用系统的门槛。", 29, VIOLET, True)
    return img


def slide_training_records():
    img = bg()
    title(img, "新增功能：训练记录", "每一次训练都有结果可查，方便复盘、对比和继续改进", 14)
    d = shadow_card(img, (110, 200, 705, 860))
    text(d, (170, 258), "训练记录列表", 38, INK, True)
    for i, item in enumerate(["深蹲训练", "开合跳训练", "弓步蹲训练"]):
        y = 350 + i * 150
        d.rounded_rectangle((170, y, 640, y + 100), radius=28, fill=(248, 251, 255, 255), outline=(226, 236, 251, 255))
        color = [VIOLET, BLUE, GREEN][i]
        d.ellipse((205, y + 30, 245, y + 70), fill=color)
        text(d, (270, y + 24), item, 29, INK, True)
        text(d, (270, y + 65), "次数 / 评分 / 问题概览", 21, MUTED)

    d = shadow_card(img, (770, 200, 1245, 470))
    text(d, (830, 260), "本次结果", 36, INK, True)
    text(d, (830, 330), "23 次", 54, VIOLET, True)
    text(d, (835, 395), "完成次数", 22, MUTED)
    text(d, (1030, 330), "3 次", 54, CORAL, True)
    text(d, (1035, 395), "需调整", 22, MUTED)

    d = shadow_card(img, (770, 535, 1245, 860))
    text(d, (830, 590), "历史趋势", 36, INK, True)
    for i, h in enumerate([60, 105, 82, 145, 122]):
        x = 850 + i * 62
        d.rounded_rectangle((x, 785 - h, x + 32, 785), radius=16, fill=[BLUE, VIOLET, GREEN, AMBER, VIOLET][i])
    d.line((820, 790, 1160, 790), fill=(216, 226, 245), width=3)

    info = [
        ("记录保存", "保存训练时间、动作类型、次数、评分、骨架数据和问题结果。", BLUE),
        ("问题复盘", "点击某次训练后，可以回看完整训练和每一次动作的纠错信息。", VIOLET),
        ("进步可见", "用历史记录呈现训练变化，让用户知道自己是否练得更稳定。", GREEN),
    ]
    for i, (h, b, c) in enumerate(info):
        y = 210 + i * 210
        d = shadow_card(img, (1310, y, 1810, y + 150), 34)
        d.rectangle((1338, y + 28, 1350, y + 122), fill=c)
        text(d, (1385, y + 30), h, 31, INK, True)
        multiline(d, (1385, y + 77), b, 23, MUTED, width=355, gap=7)
    return img


def slide_mini_program():
    img = bg()
    title(img, "新增功能：小程序端", "用手机完成动作采集，把训练数据同步到 Web 端复盘", 15)
    d = shadow_card(img, (130, 190, 575, 900), 48)
    d.rounded_rectangle((175, 235, 530, 855), radius=52, fill=(249, 252, 255, 255), outline=(220, 231, 248, 255), width=3)
    text(d, (250, 290), "姿态采集", 34, INK, True)
    d.rounded_rectangle((225, 365, 480, 590), radius=30, fill=(9, 17, 36, 245))
    paste_center(img, load_asset("start-phone-camera.png", (230, 230)), (235, 355, 470, 600))
    d.rounded_rectangle((225, 655, 480, 725), radius=35, fill=VIOLET)
    text(d, (352, 677), "开始训练", 28, (255, 255, 255), True, "ma")
    text(d, (352, 775), "实时采集骨架数据", 23, MUTED, False, "ma")

    info = [
        ("手机采集更低门槛", "用户不用额外设备，只需要打开小程序和摄像头，就能进入训练。", BLUE),
        ("动作计数与上传", "小程序负责采集关键帧与训练结果，并把骨架数据传回后端。", VIOLET),
        ("Web 端继续复盘", "训练完成后，可以在 Web 端查看 3D 回放、单次纠错和训练记录。", GREEN),
    ]
    for i, (h, b, c) in enumerate(info):
        y = 205 + i * 205
        d = shadow_card(img, (670, y, 1205, y + 145), 34)
        d.rectangle((700, y + 26, 713, y + 119), fill=c)
        text(d, (750, y + 28), h, 31, INK, True)
        multiline(d, (750, y + 75), b, 23, MUTED, width=385, gap=7)

    d = shadow_card(img, (1300, 205, 1810, 790), 38)
    text(d, (1360, 260), "端到端流程", 36, INK, True)
    steps = [("拍摄", "摄像头采集动作"), ("识别", "形成姿态关键点"), ("上传", "同步训练数据"), ("复盘", "Web 查看结果")]
    for i, (h, b) in enumerate(steps):
        y = 350 + i * 92
        color = [BLUE, VIOLET, GREEN, AMBER][i]
        d.ellipse((1365, y, 1408, y + 43), fill=color)
        if i < 3:
            d.line((1386, y + 43, 1386, y + 88), fill=(216, 226, 245), width=3)
        text(d, (1440, y - 2), h, 27, INK, True)
        text(d, (1525, y + 1), b, 22, MUTED)
    d = shadow_card(img, (300, 925, 1620, 1000), 34, (255, 255, 255, 230))
    text(d, (350, 946), "页面价值：小程序负责“随时开始练”，Web 端负责“看得清、复盘深”，形成多端联动体验。", 28, VIOLET, True)
    return img


def make_images():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    slides = [slide_action_library(), slide_training_records(), slide_mini_program()]
    paths = []
    for i, img in enumerate(slides, 13):
        path = ASSET_DIR / f"slide_{i:02d}.png"
        img.convert("RGB").save(path, quality=96)
        paths.append(path)
    sheet = Image.new("RGB", (640 * 3, 360), "white")
    for i, path in enumerate(paths):
        thumb = Image.open(path).resize((640, 360), Image.Resampling.LANCZOS)
        sheet.paste(thumb, (i * 640, 0))
    sheet.save(ASSET_DIR / "contact_sheet.jpg", quality=92)
    return paths


def move_last_slides_after(prs, count, after_index_zero_based):
    sld_id_lst = prs.slides._sldIdLst
    moving = list(sld_id_lst)[-count:]
    for el in moving:
        sld_id_lst.remove(el)
    insert_at = after_index_zero_based + 1
    for offset, el in enumerate(moving):
        sld_id_lst.insert(insert_at + offset, el)


def update_footers(prs):
    for idx, slide in enumerate(prs.slides, 1):
        if idx == 1 or idx == len(prs.slides):
            continue
        wanted = f"{idx:02d} / 姿态棱镜课程汇报"
        for shape in slide.shapes:
            if hasattr(shape, "text") and re.match(r"^\d{2}\s*/", shape.text.strip()):
                shape.text = wanted
                for p in shape.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.RIGHT
                    for run in p.runs:
                        run.font.name = "Microsoft YaHei"
                        run.font.size = Pt(8.2)
                        run.font.color.rgb = rgb("98A2B3")
                break


def make_ppt(paths):
    prs = Presentation(str(PPT_IN))
    for path in paths:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(path), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
    move_last_slides_after(prs, 3, 11)
    update_footers(prs)
    prs.save(str(PPT_OUT))


def main():
    paths = make_images()
    make_ppt(paths)
    print(PPT_OUT)
    print(ASSET_DIR / "contact_sheet.jpg")


if __name__ == "__main__":
    main()
