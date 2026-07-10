from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(r"D:\02-学习\项目实训\4")
OUT = ROOT / "deliverables"
PPT_IN = OUT / "姿态棱镜_课程项目汇报.pptx"
PPT_OUT = OUT / "姿态棱镜_课程项目汇报_新增三页.pptx"
ASSETS = ROOT / "frontend" / "src" / "assets"
LOGO = Path(r"D:\02-学习\项目实训\logo\938b2dd0fea680c4f870aa5bca309594.png")

W, H = 13.333, 7.5


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def add_text(slide, text, x, y, w, h, size=18, color="111827", bold=False, align="left", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rect(slide, x, y, w, h, fill="FFFFFF", line="E6ECF8", radius=True, shadow=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if shadow:
        try:
            shape.shadow.inherit = False
            shape.shadow.blur_radius = Pt(7)
            shape.shadow.distance = Pt(2)
            shape.shadow.angle = 45
            shape.shadow.fore_color.rgb = rgb("000000")
            shape.shadow.transparency = 86
        except Exception:
            pass
    return shape


def circle(slide, x, y, size, fill):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color="6C3BFF", width=2):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def pill(slide, text, x, y, fill="F1ECFF", color="6C3BFF", w=None):
    width = w or max(1.0, 0.16 * len(text) + 0.5)
    rect(slide, x, y, width, 0.36, fill, None, True)
    add_text(slide, text, x, y + 0.085, width, 0.16, 9.5, color, True, "center")
    return width


def title(slide, main, sub, page):
    add_text(slide, main, 0.58, 0.34, 8.8, 0.55, 29, "111827", True)
    add_text(slide, sub, 0.6, 0.92, 8.4, 0.28, 10.8, "667085")
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(11.85), Inches(0.28), width=Inches(0.56), height=Inches(0.56))
    footer(slide, page)


def footer(slide, page):
    add_text(slide, f"{page:02d} / 姿态棱镜课程汇报", 10.55, 7.12, 2.15, 0.2, 8.2, "98A2B3", False, "right")


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("F6F8FF")
    circle(slide, 11.0, -0.75, 2.4, "EAF1FF")
    circle(slide, -0.8, 5.6, 2.2, "E8F6FF")
    line(slide, 0.25, 6.72, 12.7, 5.82, "DDE8FF", 1.4)
    line(slide, 0.25, 6.92, 12.7, 6.08, "E9EFFD", 1.2)


def add_card(slide, heading, body, x, y, w, h, accent="6C3BFF"):
    rect(slide, x, y, w, h, "FFFFFF", "E0E8F7", True, True)
    rect(slide, x + 0.18, y + 0.22, 0.08, h - 0.44, accent, None, False)
    add_text(slide, heading, x + 0.42, y + 0.22, w - 0.65, 0.24, 14.5, "111827", True)
    add_text(slide, body, x + 0.42, y + 0.68, w - 0.65, h - 0.82, 10.8, "475467")


def add_bullets(slide, items, x, y, w, h, size=12, color="344054"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(7)
    return box


def image(slide, path, x, y, w, h):
    if Path(path).exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def phone_frame(slide, x, y, w, h):
    rect(slide, x, y, w, h, "FFFFFF", "DDE7F8", True, True)
    rect(slide, x + 0.18, y + 0.22, w - 0.36, h - 0.44, "F7FAFF", "ECF1FB", True)
    circle(slide, x + w / 2 - 0.14, y + 0.08, 0.28, "DDE7F8")


def action_library_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "新增功能：动作库与动作详情跟练", "用户先选动作，再看详情，最后进入跟练或训练采集", 13)

    rect(slide, 0.78, 1.32, 6.1, 4.95, "FFFFFF", "E0E8F7", True, True)
    add_text(slide, "动作库", 1.18, 1.62, 1.7, 0.3, 18, "111827", True)
    rect(slide, 1.18, 2.08, 5.2, 0.42, "F4F7FF", "E4EBF8", True)
    add_text(slide, "搜索动作 / 按部位筛选", 1.42, 2.2, 2.5, 0.16, 8.6, "667085")
    for i, name in enumerate(["深蹲", "开合跳", "平板支撑", "弓步蹲"]):
        x = 1.18 + (i % 2) * 2.65
        y = 2.82 + (i // 2) * 1.2
        rect(slide, x, y, 2.25, 0.9, "F8FAFF", "E6ECF8", True)
        circle(slide, x + 0.18, y + 0.22, 0.46, ["43B7FF", "6C3BFF", "44D989", "FFB84D"][i])
        add_text(slide, name, x + 0.78, y + 0.2, 1.2, 0.22, 12.5, "111827", True)
        add_text(slide, "适合日常训练", x + 0.78, y + 0.52, 1.1, 0.16, 8.4, "667085")

    rect(slide, 7.25, 1.32, 5.15, 4.95, "FFFFFF", "E0E8F7", True, True)
    add_text(slide, "动作详情 + 跟练", 7.7, 1.62, 2.8, 0.3, 18, "111827", True)
    rect(slide, 7.7, 2.05, 2.15, 1.5, "07111F", None, True)
    image(slide, ASSETS / "home-exercise-bird-dog.png", 9.95, 1.92, 1.8, 1.65)
    pill(slide, "难度", 7.72, 3.82, "EAF1FF", "43B7FF", 0.85)
    pill(slide, "目标部位", 8.75, 3.82, "F1ECFF", "6C3BFF", 1.25)
    pill(slide, "动作要点", 10.22, 3.82, "EAFBF3", "22A66F", 1.25)
    rect(slide, 7.7, 4.55, 1.85, 0.5, "6C3BFF", None, True)
    add_text(slide, "开始跟练", 8.08, 4.72, 1.1, 0.14, 10, "FFFFFF", True, "center")
    add_bullets(
        slide,
        ["动作库帮助用户快速找到训练项目", "详情页展示动作说明、注意点和训练入口", "跟练页承接摄像头采集，形成完整训练闭环"],
        7.7,
        5.28,
        4.25,
        0.78,
        9.6,
    )

    add_card(slide, "页面价值", "把“找动作、学动作、开始练”连在一起，降低普通用户第一次使用系统的门槛。", 1.1, 6.42, 10.9, 0.55, "6C3BFF")
    return slide


def training_record_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "新增功能：训练记录", "每一次训练都有结果可查，方便复盘、对比和继续改进", 14)

    rect(slide, 0.78, 1.35, 4.0, 4.88, "FFFFFF", "E0E8F7", True, True)
    add_text(slide, "训练记录列表", 1.16, 1.72, 2.0, 0.28, 17, "111827", True)
    for i, item in enumerate(["深蹲训练", "开合跳训练", "弓步蹲训练"]):
        y = 2.32 + i * 1.05
        rect(slide, 1.1, y, 3.25, 0.72, "F8FAFF", "E6ECF8", True)
        circle(slide, 1.32, y + 0.2, 0.32, ["6C3BFF", "43B7FF", "44D989"][i])
        add_text(slide, item, 1.82, y + 0.17, 1.35, 0.18, 11.2, "111827", True)
        add_text(slide, "次数 / 评分 / 问题概览", 1.82, y + 0.43, 1.75, 0.14, 8.4, "667085")

    rect(slide, 5.1, 1.35, 3.35, 2.05, "FFFFFF", "E0E8F7", True, True)
    add_text(slide, "本次结果", 5.48, 1.7, 1.5, 0.25, 16, "111827", True)
    add_text(slide, "23 次", 5.48, 2.12, 1.1, 0.38, 26, "6C3BFF", True)
    add_text(slide, "完成次数", 5.55, 2.58, 1.2, 0.16, 8.6, "667085")
    add_text(slide, "3 次", 6.95, 2.12, 1.0, 0.38, 26, "FF5B57", True)
    add_text(slide, "需调整", 6.98, 2.58, 1.0, 0.16, 8.6, "667085")

    rect(slide, 5.1, 3.75, 3.35, 2.48, "FFFFFF", "E0E8F7", True, True)
    add_text(slide, "历史趋势", 5.48, 4.1, 1.5, 0.25, 16, "111827", True)
    for i, h in enumerate([0.32, 0.55, 0.42, 0.72, 0.62]):
        rect(slide, 5.55 + i * 0.46, 5.45 - h, 0.22, h, ["43B7FF", "6C3BFF", "44D989", "FFB84D", "6C3BFF"][i], None, True)
    line(slide, 5.48, 5.48, 7.95, 5.48, "D8E2F5", 1.2)

    add_card(slide, "记录保存", "保存训练时间、动作类型、次数、评分、骨架数据和问题结果。", 8.75, 1.35, 3.75, 1.18, "43B7FF")
    add_card(slide, "问题复盘", "点击某次训练后，可以回看完整训练和每一次动作的纠错信息。", 8.75, 2.9, 3.75, 1.18, "6C3BFF")
    add_card(slide, "进步可见", "用历史记录呈现训练变化，让用户知道自己是否练得更稳定。", 8.75, 4.45, 3.75, 1.18, "44D989")
    return slide


def mini_program_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "新增功能：小程序端", "用手机完成动作采集，把训练数据同步到 Web 端复盘", 15)

    phone_frame(slide, 0.95, 1.35, 3.1, 5.25)
    add_text(slide, "姿态采集", 1.42, 1.88, 1.25, 0.22, 14, "111827", True)
    rect(slide, 1.35, 2.35, 2.3, 2.1, "07111F", None, True)
    image(slide, ASSETS / "start-phone-camera.png", 1.62, 2.55, 1.7, 1.7)
    rect(slide, 1.35, 4.75, 2.3, 0.42, "6C3BFF", None, True)
    add_text(slide, "开始训练", 1.94, 4.89, 1.1, 0.14, 9.5, "FFFFFF", True, "center")
    add_text(slide, "实时采集骨架数据", 1.55, 5.48, 1.7, 0.16, 8.8, "667085", False, "center")

    add_card(slide, "手机采集更低门槛", "用户不用额外设备，只需要打开小程序和摄像头，就能进入训练。", 4.55, 1.45, 3.55, 1.25, "43B7FF")
    add_card(slide, "动作计数与上传", "小程序负责采集关键帧与训练结果，并把骨架数据传回后端。", 4.55, 3.0, 3.55, 1.25, "6C3BFF")
    add_card(slide, "Web 端继续复盘", "训练完成后，可以在 Web 端查看 3D 回放、单次纠错和训练记录。", 4.55, 4.55, 3.55, 1.25, "44D989")

    rect(slide, 8.55, 1.45, 3.8, 4.35, "FFFFFF", "E0E8F7", True, True)
    add_text(slide, "端到端流程", 8.95, 1.82, 1.7, 0.26, 17, "111827", True)
    steps = [("拍摄", "摄像头采集动作"), ("识别", "形成姿态关键点"), ("上传", "同步训练数据"), ("复盘", "Web 查看结果")]
    for i, (head, body) in enumerate(steps):
        y = 2.42 + i * 0.78
        circle(slide, 8.98, y, 0.32, ["43B7FF", "6C3BFF", "44D989", "FFB84D"][i])
        if i < 3:
            line(slide, 9.14, y + 0.33, 9.14, y + 0.73, "D8E2F5", 1.2)
        add_text(slide, head, 9.46, y - 0.02, 0.75, 0.18, 11.2, "111827", True)
        add_text(slide, body, 10.08, y - 0.01, 1.7, 0.16, 9, "667085")

    add_card(slide, "页面价值", "小程序负责“随时开始练”，Web 端负责“看得清、复盘深”，形成多端联动体验。", 1.1, 6.58, 10.9, 0.5, "6C3BFF")
    return slide


def move_last_slides_after(prs, count, after_index_zero_based):
    sld_id_lst = prs.slides._sldIdLst
    moving = list(sld_id_lst)[-count:]
    for el in moving:
        sld_id_lst.remove(el)
    insert_at = after_index_zero_based + 1
    for offset, el in enumerate(moving):
        sld_id_lst.insert(insert_at + offset, el)


def update_footers(prs):
    for idx, slide in enumerate(prs.slides, 1):
        if idx == 1 or idx == len(prs.slides):
            continue
        wanted = f"{idx:02d} / 姿态棱镜课程汇报"
        found = False
        for shape in slide.shapes:
            if hasattr(shape, "text") and re.match(r"^\d{2}\s*/", shape.text.strip()):
                shape.text = wanted
                for p in shape.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.RIGHT
                    for run in p.runs:
                        run.font.name = "Microsoft YaHei"
                        run.font.size = Pt(8.2)
                        run.font.color.rgb = rgb("98A2B3")
                found = True
        if not found and idx != len(prs.slides):
            footer(slide, idx)


def main():
    prs = Presentation(str(PPT_IN))
    action_library_slide(prs)
    training_record_slide(prs)
    mini_program_slide(prs)
    move_last_slides_after(prs, 3, 11)
    update_footers(prs)
    prs.save(str(PPT_OUT))
    print(PPT_OUT)


if __name__ == "__main__":
    main()
